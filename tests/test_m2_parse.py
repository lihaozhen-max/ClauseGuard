"""M2 解析与字段提取用例（AC04、AC06、AC07、TS-03、TS-06、TS-07）。

分两部分：
- **单元部分**直接构造 ``ParsedDocument``，不依赖数据库与 OCR，覆盖字段抽取规则与边界；
- **集成部分**跑真实的"下载 + 解析"链路，验证 AC04/AC06/AC07 的判据。
"""

from __future__ import annotations

import re
from pathlib import Path

import pytest
from sqlalchemy import func, select, text

from app.core.enums import ExtractStatus, ParseMode
from app.core.errors import AppError, ErrorCode
from app.db.models import ApprovalTask, ContractParse
from app.db.session import session_scope
from app.modules.parser.fields import (
    ALL_FIELDS,
    BASIC_FIELDS,
    CLAUSE_FIELDS,
    chinese_number_to_int,
    clause_heading_index,
    extract_fields,
)
from app.modules.parser.cleaning import build_document
from app.modules.parser.extractors import extract_document
from app.modules.parser.models import ParsedDocument, TextBlock
from app.tools.attachment import download_contract_attachment
from app.tools.parser import parse_contract_document, parse_task

POSITION_TEXT_RE = re.compile(r"^第\d+页 第\d+段$")
POSITION_OCR_RE = re.compile(r"^第\d+页 区域\(\d+,\d+\)$")


def build_doc(pages: dict[int, list[str]], parse_mode: ParseMode = ParseMode.TEXT) -> ParsedDocument:
    """用页码 → 段落列表构造 ParsedDocument（偏移与全文按真实规则拼装）。"""
    blocks: list[TextBlock] = []
    chunks: list[str] = []
    cursor = 0
    for page_no in sorted(pages):
        for para_no, para in enumerate(pages[page_no], start=1):
            blocks.append(
                TextBlock(page_no, para_no, para, cursor, cursor + len(para))
            )
            chunks.append(para)
            cursor += len(para) + 1
    return ParsedDocument(
        full_text="\n".join(chunks),
        blocks=blocks,
        parse_mode=parse_mode,
        page_count=len(pages),
    )


def statuses(pages: dict[int, list[str]]) -> dict[str, str]:
    basic, clauses, _ = extract_fields(build_doc(pages))
    return {r.field_name: r.extract_status for r in [*basic, *clauses]}


def values(pages: dict[int, list[str]]) -> dict[str, str | None]:
    basic, clauses, _ = extract_fields(build_doc(pages))
    return {r.field_name: r.field_value for r in [*basic, *clauses]}


# ── 金额 ──────────────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    ("line", "expected"),
    [
        ("本合同总金额为人民币伍拾万元整（¥500000.00）。", "500000"),
        ("合同金额：人民币300000元", "300000"),
        ("金额 ¥1,234,567.00 元", "1234567"),
        ("本合同总金额为人民币伍拾万元整。", "500000"),
        ("合同价款为人民币叁拾万元整。", "300000"),
        ("金额为人民币壹佰贰拾万元整。", "1200000"),
    ],
)
def test_amount_normalization(line: str, expected: str) -> None:
    """SPEC §2.3：金额必须是**纯数字字符串**（无千分位、无货币符号）。"""
    assert values({1: [line]})["contract_amount"] == expected


@pytest.mark.parametrize(
    ("cn", "expected"),
    [("伍拾万", 500000), ("叁拾万", 300000), ("壹佰贰拾万", 1200000), ("十五", 15), ("两千零五", 2005)],
)
def test_chinese_number_to_int(cn: str, expected: int) -> None:
    assert chinese_number_to_int(cn) == expected


# ── 币种 ──────────────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    ("line", "expected"),
    [
        ("币种为人民币（CNY）。", "CNY"),
        ("币种：USD", "USD"),
        ("合同以欧元结算", "EUR"),
        ("以港币支付", "HKD"),
    ],
)
def test_currency_extraction(line: str, expected: str) -> None:
    assert values({1: [line]})["currency"] == expected


# ── 日期 ──────────────────────────────────────────────────────────────────


def test_date_extraction_prefers_effective_and_expiry_context() -> None:
    doc = {2: ["本合同自2026年10月1日起生效，有效期至2027年9月30日止。"]}
    assert values(doc)["effective_date"] == "2026-10-01"
    assert values(doc)["expiry_date"] == "2027-09-30"


def test_iso_date_format_supported() -> None:
    assert values({1: ["本合同自 2026-10-01 起生效。"]})["effective_date"] == "2026-10-01"


def test_illegal_date_rejected() -> None:
    assert statuses({1: ["生效日期 2026年13月45日"]})["effective_date"] == ExtractStatus.MISSING.value


# ── 主体 ──────────────────────────────────────────────────────────────────


def test_party_extraction_skips_seal_lines() -> None:
    doc = {
        1: [
            "甲方（签约主体）：某某集团有限公司",
            "乙方（对方名称）：某某科技有限公司",
            "甲方（盖章）：某某集团有限公司",
            "乙方（盖章）：某某科技有限公司",
        ]
    }
    result = values(doc)
    assert result["party_a"] == "某某集团有限公司"
    assert result["party_b"] == "某某科技有限公司"
    basic, _, _ = extract_fields(build_doc(doc))
    by_name = {r.field_name: r for r in basic}
    assert by_name["party_a"].position == "第1页 第1段"  # 命中第一行，不是盖章行


def test_party_missing_when_absent() -> None:
    assert statuses({1: ["本合同一式两份。"]})["party_a"] == ExtractStatus.MISSING.value


# ── 标题 ──────────────────────────────────────────────────────────────────


def test_contract_title_prefers_explicit_field() -> None:
    doc = {1: ["采购合同", "合同编号：CG-1", "合同标题：服务器采购合同"]}
    assert values(doc)["contract_title"] == "服务器采购合同"


def test_contract_title_falls_back_to_first_short_line() -> None:
    assert values({1: ["技术服务合同", "合同编号：CG-2"]})["contract_title"] == "技术服务合同"


# ── 条款识别（PS-09）───────────────────────────────────────────────────────


def test_clause_list_item_is_not_a_heading() -> None:
    """付款条款里的列表项含"验收"二字，不能把验收条款定位到付款条款里（实测踩过）。"""
    doc = {
        1: [
            "第三条  付款方式",
            "1. 合同签订后十日内支付合同总额30%作为预付款。",
            "2. 服务期满并经甲方验收合格后30日内支付剩余70%款项。",
            "第五条  验收条款",
            "甲方应在收到报告后10个工作日内进行验收。",
        ]
    }
    headings = clause_heading_index(build_doc(doc))
    assert headings[0] == "payment_clause"
    assert headings[3] == "acceptance_clause"
    basic, clauses, _ = extract_fields(build_doc(doc))
    acceptance = next(r for r in clauses if r.field_name == "acceptance_clause")
    assert acceptance.position == "第1页 第4段"
    assert "第五条" in (acceptance.source_text or "")


def test_clause_missing_when_absent() -> None:
    doc = {1: ["第三条  付款方式", "合同签订后支付全款。"]}
    result = statuses(doc)
    assert result["confidentiality_clause"] == ExtractStatus.MISSING.value
    assert result["ip_clause"] == ExtractStatus.MISSING.value
    assert result["data_clause"] == ExtractStatus.MISSING.value
    assert result["payment_clause"] == ExtractStatus.SUCCESS.value


def test_clause_span_stops_at_next_heading() -> None:
    doc = {
        1: [
            "第七条  保密条款",
            "双方承担保密义务，保密期限为三年。",
            "第八条  数据条款",
            "乙方应删除全部数据。",
        ]
    }
    _, clauses, _ = extract_fields(build_doc(doc))
    confidentiality = next(r for r in clauses if r.field_name == "confidentiality_clause")
    assert "保密期限为三年" in confidentiality.source_text
    assert "数据条款" not in confidentiality.source_text


def test_strong_keyword_fallback_without_numbered_heading() -> None:
    """真实合同常不写"第 N 条"：只要出现强关键词也应识别为条款。"""
    doc = {1: ["保密", "双方对本合同内容承担保密义务。"]}
    assert statuses(doc)["confidentiality_clause"] == ExtractStatus.SUCCESS.value


# ── FieldRecord 结构与位置格式（§2.4 / §2.6）────────────────────────────────


def test_every_field_is_five_key_record() -> None:
    basic, clauses, _ = extract_fields(build_doc({1: ["合同标题：X", "合同编号：Y"]}))
    records = [*basic, *clauses]
    assert [r.field_name for r in records] == list(ALL_FIELDS)
    for record in records:
        payload = record.to_dict()
        assert list(payload) == [
            "field_name",
            "field_value",
            "source_text",
            "position",
            "extract_status",
        ]


def test_position_format_matches_spec_for_text_mode() -> None:
    doc = {1: ["合同标题：X"], 2: ["合同编号：Y"]}
    basic, _, _ = extract_fields(build_doc(doc))
    for record in basic:
        if record.extract_status == ExtractStatus.SUCCESS.value:
            assert POSITION_TEXT_RE.match(record.position or ""), record.position


def test_position_format_matches_spec_for_ocr_mode() -> None:
    block = TextBlock(1, 1, "合同标题：X", 0, 7, bbox=(120, 480))
    document = ParsedDocument(
        full_text="合同标题：X", blocks=[block], parse_mode=ParseMode.OCR, page_count=1
    )
    basic, _, _ = extract_fields(document)
    title = next(r for r in basic if r.field_name == "contract_title")
    assert title.position == "第1页 区域(120,480)"
    assert POSITION_OCR_RE.match(title.position)


def test_source_text_is_contiguous_substring() -> None:
    doc = {
        1: ["合同标题：服务器采购合同", "本合同总金额为人民币伍拾万元整（¥500000.00）。"],
        2: ["第三条  付款方式", "合同签订后三日内支付合同总额80%作为预付款。"],
    }
    document = build_doc(doc)
    basic, clauses, _ = extract_fields(document)
    for record in [*basic, *clauses]:
        if record.extract_status != ExtractStatus.SUCCESS.value:
            continue
        assert record.source_text, record.field_name
        assert record.source_text in document.full_text, record.field_name


def test_long_clause_is_truncated_with_ellipsis() -> None:
    body = "本合同双方约定如下：" + "甲乙双方应严格遵守本条款各项约定。" * 200
    doc = {1: ["第七条  保密条款", body]}
    _, clauses, _ = extract_fields(build_doc(doc))
    confidentiality = next(r for r in clauses if r.field_name == "confidentiality_clause")
    assert confidentiality.extract_status == ExtractStatus.SUCCESS.value
    assert confidentiality.source_text.endswith("…")
    assert len(confidentiality.source_text) <= 2000
    # 截断前的部分必须是原文连续子串（§2.4 与 FR-PARSE-06 的折中口径）
    assert confidentiality.source_text[:-1] in build_doc(doc).full_text


# ── 集成：AC04 / AC06 / AC07 ──────────────────────────────────────────────


async def _task_of(instance_id: str) -> ApprovalTask | None:
    async with session_scope() as session:
        task = await session.scalar(
            select(ApprovalTask).where(ApprovalTask.instance_id == instance_id)
        )
        if task is not None:
            session.expunge(task)
    return task


@pytest.fixture(scope="module", autouse=True)
def _pulled(db_runner, mock_app):
    """确保 AP-001…AP-005 都已拉取为任务。

    集成用例不能假设"别的文件先跑过、顺手把待办拉下来了"——那样
    ``pytest -k test_parse_text_pdf_end_to_end_ac04`` 之类的单点运行就会失败。
    """
    import httpx

    from app.clients.approval_client import ApprovalSystemClient

    from app.tools.approval import list_pending_contract_approvals

    client = ApprovalSystemClient(transport=httpx.ASGITransport(app=mock_app))
    db_runner(lambda: list_pending_contract_approvals(20, client=client))


@pytest.mark.requires_db
def test_parse_text_pdf_end_to_end_ac04(live_db: str, db_runner, approval_client) -> None:
    """AC04：用 AP-001（文本型 PDF）解析 → parse_status=success，16 个字段全部产出 FieldRecord。"""
    task = db_runner(lambda: _task_of("AP-001"))
    assert task is not None
    outcome = db_runner(lambda: parse_task(task.id, client=approval_client))

    assert outcome.parse_mode == ParseMode.TEXT.value
    assert outcome.parse_status == "success"
    assert outcome.page_count == 2
    assert outcome.used_ocr_pages == []
    assert len(outcome.basic_info) == 8
    assert len(outcome.clause_info) == 8
    assert [r["field_name"] for r in outcome.basic_info] == list(BASIC_FIELDS)
    assert [r["field_name"] for r in outcome.clause_info] == list(CLAUSE_FIELDS)


@pytest.mark.requires_db
def test_all_16_fields_have_value_or_explicit_missing_ac06(live_db: str, db_runner) -> None:
    """AC06：8 基本信息 + 8 条款均有值（或明确 missing）。"""

    async def fetch():
        async with session_scope() as session:
            row = await session.scalar(
                select(ContractParse)
                .join(ApprovalTask, ApprovalTask.id == ContractParse.task_id)
                .where(ApprovalTask.instance_id == "AP-001")
                .order_by(ContractParse.id.desc())
            )
            if row is not None:
                session.expunge(row)
            return row

    row = db_runner(fetch)
    assert row is not None
    records = [*row.basic_info_json, *row.clause_info_json]
    assert len(records) == 16
    for record in records:
        assert record["extract_status"] in {"success", "missing", "failed"}
        if record["extract_status"] == "success":
            assert record["field_value"]
            assert record["source_text"]
            assert record["position"]
        else:
            assert record["field_value"] is None


@pytest.mark.requires_db
def test_sampled_fields_trace_back_to_source_ac07(live_db: str, db_runner) -> None:
    """AC07：抽查字段的 source_text 是 full_text 连续子串，position 符合 §2.6 正则。"""

    async def fetch():
        async with session_scope() as session:
            row = await session.scalar(
                select(ContractParse)
                .join(ApprovalTask, ApprovalTask.id == ContractParse.task_id)
                .where(ApprovalTask.instance_id == "AP-001")
                .order_by(ContractParse.id.desc())
            )
            if row is not None:
                session.expunge(row)
            return row

    row = db_runner(fetch)
    assert row is not None and row.full_text
    by_name = {r["field_name"]: r for r in [*row.basic_info_json, *row.clause_info_json]}

    for field in ("contract_amount", "payment_clause", "effective_date"):
        record = by_name[field]
        assert record["extract_status"] == "success", field
        assert record["source_text"] in row.full_text, field
        assert POSITION_TEXT_RE.match(record["position"]), record["position"]

    # 全部 success 字段都满足同一条可追溯性约束（比 AC07 的抽查更严）
    for record in by_name.values():
        if record["extract_status"] == "success":
            assert record["source_text"] in row.full_text, record["field_name"]
            assert POSITION_TEXT_RE.match(record["position"]), record["field_name"]


@pytest.mark.requires_db
def test_reparse_updates_same_document_fr_parse_09(live_db: str, db_runner, approval_client) -> None:
    """FR-PARSE-09：重复解析同一附件必须更新原记录，不新增。"""
    task = db_runner(lambda: _task_of("AP-001"))
    assert task is not None

    first = db_runner(lambda: parse_task(task.id, client=approval_client))
    second = db_runner(lambda: parse_contract_document(first.document_id))

    assert second.document_id == first.document_id

    async def count():
        async with session_scope() as session:
            return int(
                await session.scalar(
                    select(func.count())
                    .select_from(ContractParse)
                    .where(ContractParse.task_id == task.id)
                )
                or 0
            )

    assert db_runner(count) == 1


@pytest.mark.requires_db
def test_empty_attachment_blocks_with_parse_error_ts06(live_db: str, db_runner, approval_client) -> None:
    """TS-06 / FR-PARSE-07/08：0 字节文件 → blocked + parse_error，禁止只返回空结果。"""
    task = db_runner(lambda: _task_of("AP-005"))
    assert task is not None

    with pytest.raises(AppError) as excinfo:
        db_runner(lambda: parse_task(task.id, client=approval_client))
    assert excinfo.value.code is ErrorCode.EMPTY_CONTRACT_CONTENT

    async def fetch():
        async with session_scope() as session:
            task_row = (
                await session.execute(
                    text(
                        "SELECT task_status, blocked_stage, error_code FROM approval_tasks WHERE id=:t"
                    ),
                    {"t": task.id},
                )
            ).first()
            parse_row = await session.scalar(
                select(ContractParse).where(ContractParse.task_id == task.id)
            )
            if parse_row is not None:
                session.expunge(parse_row)
        return dict(task_row._mapping), parse_row

    state, parse_row = db_runner(fetch)
    assert state["task_status"] == "blocked"
    assert state["blocked_stage"] == "parsing"
    assert state["error_code"] == ErrorCode.EMPTY_CONTRACT_CONTENT.value
    assert parse_row is not None
    assert parse_row.parse_status == "failed"
    assert parse_row.parse_error and "EMPTY_CONTRACT_CONTENT" in parse_row.parse_error


@pytest.mark.requires_db
def test_parse_document_id_falls_back_to_attachment_id(
    live_db: str, db_runner, approval_client
) -> None:
    """document_id 口径：先按 contract_parses.id，再按 approval_attachments.id 兜底。"""
    # 先确保 AP-001 的附件已下载（用例自足，不依赖其它用例的执行顺序）
    db_runner(
        lambda: download_contract_attachment(
            "AP-001", "ATT-001", "AP-001_采购合同.pdf", client=approval_client
        )
    )

    async def attachment_id_of(instance_id: str) -> int:
        async with session_scope() as session:
            row = (
                await session.execute(
                    text(
                        "SELECT a.id FROM approval_attachments a "
                        "JOIN approval_tasks t ON t.id = a.task_id WHERE t.instance_id = :i "
                        "ORDER BY a.id LIMIT 1"
                    ),
                    {"i": instance_id},
                )
            ).first()
        assert row is not None, f"{instance_id} 尚无附件记录"
        return int(row.id)

    attachment_pk = db_runner(lambda: attachment_id_of("AP-001"))
    outcome = db_runner(lambda: parse_contract_document(attachment_pk))
    assert outcome.task_id > 0
    assert outcome.parse_status == "success"


# ── FR-PARSE-01：Word 文档 ────────────────────────────────────────────────


def test_docx_extraction_full_pipeline(tmp_path: Path) -> None:
    """FR-PARSE-01：``.docx`` 必须支持；段落即"段"，position 走 text 格式。"""
    import docx

    document = docx.Document()
    for line in [
        "技术服务合同",
        "合同编号：CG-DOCX-2026-0009",
        "合同标题：文档解析测试合同",
        "甲方（签约主体）：某某集团有限公司",
        "乙方（对方名称）：某某软件服务有限公司",
        "第二条  合同金额与币种",
        "本合同总金额为人民币叁拾万元整（¥300000.00）。",
        "币种为人民币（CNY）。",
        "第四条  服务期限",
        "本合同自2026年10月1日起生效，有效期至2027年9月30日止。",
        "第七条  保密条款",
        "双方对商业秘密承担保密义务，保密期限为三年。",
    ]:
        document.add_paragraph(line)
    path = tmp_path / "sample.docx"
    document.save(path)

    extraction = extract_document(path, "sample.docx")
    assert extraction.kind.value == "docx"
    assert extraction.parse_mode is ParseMode.TEXT

    parsed = build_document(
        extraction.blocks, extraction.parse_mode, extraction.page_count
    )
    basic, clauses, _ = extract_fields(parsed)
    by_name = {r.field_name: r for r in [*basic, *clauses]}

    assert by_name["contract_no"].field_value == "CG-DOCX-2026-0009"
    assert by_name["contract_amount"].field_value == "300000"
    assert by_name["currency"].field_value == "CNY"
    assert by_name["effective_date"].field_value == "2026-10-01"
    assert by_name["expiry_date"].field_value == "2027-09-30"
    assert by_name["confidentiality_clause"].extract_status == "success"
    assert by_name["payment_clause"].extract_status == "missing"  # 文档里确实没有付款条款
    assert by_name["contract_title"].position == "第1页 第3段"
